"""
Leitura do export do MacroFactor (.xlsx).

O MacroFactor não tem API: o caminho estável é o export manual (More → Data
Management → Data Export), uma planilha com uma aba por assunto. Daqui saem as
abas que fecham a conta do peso com o treino:

- `Calories & Macros` e `Micronutrients` — o que entrou, por dia registrado
- `Expenditure` — gasto total estimado pelo app, por dia
- `Weight Trend` — peso suavizado, por dia
- `Scale Weight` — pesagens, com percentual de gordura
- `Nutrition Program Settings` — metas de calorias e macros por dia da semana,
  uma leva a cada revisão do programa

`parse_export` só lê e normaliza; `import_export` grava no catálogo, e é o que
`scripts/import_macrofactor.py` chama. As abas de musculação, receitas e
configurações ficam de fora.
"""
from __future__ import annotations

from bisect import bisect_right
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from storage.catalog_store import CatalogStore

WEEKDAYS = (
    "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday",
)


@dataclass
class NutritionDay:
    """Um dia do export, com o que houver: ingestão, gasto, tendência, metas."""

    date: date
    calories_kcal: float | None = None
    protein_g: float | None = None
    carbs_g: float | None = None
    fat_g: float | None = None
    fiber_g: float | None = None
    expenditure_kcal: float | None = None
    trend_weight_kg: float | None = None
    target_calories_kcal: float | None = None
    target_protein_g: float | None = None
    target_carbs_g: float | None = None
    target_fat_g: float | None = None

    def fields(self) -> dict[str, float | None]:
        """Colunas de `nutrition_daily`, sem a data."""
        return {k: v for k, v in self.__dict__.items() if k != "date"}


@dataclass
class WeighIn:
    """Uma pesagem na balança."""

    date: date
    weight_kg: float
    body_fat_pct: float | None = None


@dataclass
class MacroFactorExport:
    """Conteúdo aproveitado de um export."""

    days: list[NutritionDay] = field(default_factory=list)
    weigh_ins: list[WeighIn] = field(default_factory=list)


def parse_export(path: str | Path) -> MacroFactorExport:
    """Lê o export do MacroFactor e devolve os dias e as pesagens.

    Args:
        path: caminho do `.xlsx` exportado pelo app.

    Returns:
        Dias em ordem cronológica (a união das datas de ingestão, gasto e
        tendência) e as pesagens.

    Raises:
        ValueError: se o arquivo não tiver nenhuma das abas esperadas.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {ws.title: list(ws.iter_rows(values_only=True)) for ws in workbook}
    finally:
        workbook.close()

    if not {"Calories & Macros", "Expenditure", "Scale Weight"} & set(sheets):
        raise ValueError(f"{path} não parece um export do MacroFactor")

    days: dict[date, NutritionDay] = {}

    def day(when: date) -> NutritionDay:
        return days.setdefault(when, NutritionDay(date=when))

    for record in _records(sheets.get("Calories & Macros")):
        calories = _number(record.get("Calories (kcal)"))
        if not calories:
            # O app exporta linha zerada para dia aberto sem comida anotada
            # (só água ou suplemento). Zero aqui é "não registrou".
            continue
        entry = day(record["date"])
        entry.calories_kcal = calories
        entry.protein_g = _number(record.get("Protein (g)"))
        entry.carbs_g = _number(record.get("Carbs (g)"))
        entry.fat_g = _number(record.get("Fat (g)"))

    for record in _records(sheets.get("Micronutrients")):
        fiber = _number(record.get("Fiber (g)"))
        if fiber is not None:
            day(record["date"]).fiber_g = fiber

    for record in _records(sheets.get("Expenditure")):
        value = _number(record.get("Expenditure"))
        if value is not None:
            day(record["date"]).expenditure_kcal = value

    for record in _records(sheets.get("Weight Trend")):
        value = _number(record.get("Trend Weight (kg)"))
        if value is not None:
            day(record["date"]).trend_weight_kg = value

    _apply_targets(days, sheets.get("Nutrition Program Settings"))

    weigh_ins = [
        WeighIn(
            date=record["date"],
            weight_kg=weight,
            body_fat_pct=_number(record.get("Fat Percent")),
        )
        for record in _records(sheets.get("Scale Weight"))
        if (weight := _number(record.get("Weight (kg)"))) is not None
    ]

    return MacroFactorExport(
        days=[days[when] for when in sorted(days)],
        weigh_ins=sorted(weigh_ins, key=lambda w: w.date),
    )


def import_export(
    store: CatalogStore, athlete_id: str, export: MacroFactorExport
) -> dict[str, int]:
    """Grava os dias e as pesagens do export.

    Idempotente: reimportar atualiza os dias em vez de duplicar. As pesagens só
    preenchem o que o Garmin não registrou (`CatalogStore.upsert_body_weight`).

    Args:
        store: catálogo aberto para escrita.
        athlete_id: UUID do atleta.
        export: conteúdo lido por `parse_export`.

    Returns:
        Dias gravados e pesagens que acrescentaram algum valor.
    """
    for day in export.days:
        store.upsert_nutrition_daily(
            athlete_id, day.date, source="macrofactor", **day.fields()
        )
    weigh_ins = sum(
        store.upsert_body_weight(
            athlete_id, w.date, w.weight_kg, body_fat_pct=w.body_fat_pct
        )
        for w in export.weigh_ins
    )
    return {"days": len(export.days), "weigh_ins": weigh_ins}


def _apply_targets(days: dict[date, NutritionDay], rows: list[tuple] | None) -> None:
    """Preenche as metas de cada dia com o programa vigente naquela data.

    O programa é revisto de tempos em tempos, e cada revisão traz sete linhas
    (uma por dia da semana). Vale para um dia a revisão mais recente que não
    seja posterior a ele; dias anteriores à primeira revisão ficam sem meta.
    """
    programs: dict[date, dict[str, dict[str, Any]]] = {}
    for record in _records(rows, date_column="Program Update Date"):
        weekday = record.get("Program Weekday")
        if weekday in WEEKDAYS:
            programs.setdefault(record["date"], {})[weekday] = record

    updates = sorted(programs)
    for when, entry in days.items():
        position = bisect_right(updates, when)
        if not position:
            continue
        target = programs[updates[position - 1]].get(WEEKDAYS[when.weekday()])
        if not target:
            continue
        entry.target_calories_kcal = _number(target.get("Calories (kcal)"))
        entry.target_protein_g = _number(target.get("Protein (g)"))
        entry.target_carbs_g = _number(target.get("Carbs (g)"))
        entry.target_fat_g = _number(target.get("Fat (g)"))


def _records(
    rows: list[tuple] | None, date_column: str = "Date"
) -> list[dict[str, Any]]:
    """Linhas de uma aba como dicionários, só as que têm data legível."""
    if not rows:
        return []
    header = [str(cell) if cell is not None else "" for cell in rows[0]]
    if date_column not in header:
        return []
    records = []
    for values in rows[1:]:
        record = dict(zip(header, values))
        when = _as_date(record.get(date_column))
        if when is not None:
            record["date"] = when
            records.append(record)
    return records


def _as_date(value: Any) -> date | None:
    """Data da célula: o export mistura datetime e texto `dd/mm/aaaa`."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                continue
    return None


def _number(value: Any) -> float | None:
    """Número da célula, ou None para vazio e texto."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None
