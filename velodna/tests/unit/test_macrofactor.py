"""Leitura do export do MacroFactor e gravação em nutrition_daily.

A planilha de teste é montada aqui, com valores inventados: o export real traz
dados de saúde e fica em `docs/dados-saude/`, fora do git.
"""
from __future__ import annotations

from datetime import date, datetime

import duckdb
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from api.dependencies import get_db
from api.main import app
from ingestion.macrofactor import import_export, parse_export
from storage.catalog_store import CatalogStore

ATHLETE_ID = "11111111-1111-1111-1111-111111111111"


def _sheet(workbook: Workbook, title: str, rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    for row in rows:
        sheet.append(row)


@pytest.fixture
def export_path(tmp_path):
    """Export mínimo: dois dias com comida, um só com gasto e tendência."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Calories & Macros", [
        ["Date", "Calories (kcal)", "Fat (g)", "Carbs (g)", "Protein (g)"],
        [datetime(2026, 3, 2), 2000.0, 60.0, 250.0, 120.0],   # segunda
        [datetime(2026, 3, 10), 1800.0, 50.0, 200.0, 140.0],  # terça
        [datetime(2026, 3, 11), 0.0, 0.0, 0.0, 0.0],          # aberto, sem comida
    ])
    _sheet(workbook, "Micronutrients", [
        ["Date", "Alcohol (g)", "Fiber (g)"],
        [datetime(2026, 3, 2), None, 28.0],
        [datetime(2026, 3, 10), None, None],
    ])
    _sheet(workbook, "Expenditure", [
        ["Date", "Expenditure"],
        [datetime(2026, 3, 2), 2500.0],
        [datetime(2026, 3, 3), 2480.0],
        [datetime(2026, 3, 10), 2450.0],
    ])
    _sheet(workbook, "Weight Trend", [
        ["Date", "Trend Weight (kg)"],
        [datetime(2026, 3, 2), 74.2],
        [datetime(2026, 3, 3), 74.1],
    ])
    _sheet(workbook, "Scale Weight", [
        ["Date", "Weight (kg)", "Fat Percent"],
        [datetime(2026, 3, 2), 74.5, 15.0],
        [datetime(2026, 3, 3), 74.0, None],
    ])
    # Duas revisões do programa; a segunda só vale a partir de 09/03.
    _sheet(workbook, "Nutrition Program Settings", [
        ["Program Update Date", "Program Weekday", "Calories (kcal)", "Fat (g)",
         "Protein (g)", "Carbs (g)"],
        ["02/03/2026", "Monday", 2200.0, 60.0, 150.0, 260.0],
        ["02/03/2026", "Tuesday", 2100.0, 58.0, 150.0, 240.0],
        ["09/03/2026", "Monday", 1900.0, 55.0, 155.0, 200.0],
        ["09/03/2026", "Tuesday", 1850.0, 54.0, 155.0, 190.0],
    ])
    _sheet(workbook, "Recipes", [["Recipe Name"], [""]])
    path = tmp_path / "MacroFactor-20260311000000.xlsx"
    workbook.save(path)
    return path


@pytest.fixture
def db():
    conn = duckdb.connect(":memory:")
    CatalogStore(conn).initialize_schema()
    conn.execute(
        "INSERT INTO athletes (id, name, ftp_w) VALUES (?, ?, ?)",
        [ATHLETE_ID, "Teste", 250.0],
    )
    yield conn
    conn.close()


@pytest.fixture
def client(db):
    app.dependency_overrides[get_db] = lambda: db
    yield TestClient(app)
    app.dependency_overrides.clear()


# ---------------------------------------------------------------------------
# Leitura
# ---------------------------------------------------------------------------


def test_parse_merges_sheets_by_date(export_path):
    export = parse_export(export_path)
    by_date = {d.date: d for d in export.days}

    assert sorted(by_date) == [date(2026, 3, 2), date(2026, 3, 3), date(2026, 3, 10)]
    monday = by_date[date(2026, 3, 2)]
    assert (monday.calories_kcal, monday.protein_g, monday.carbs_g, monday.fat_g) == (
        2000.0, 120.0, 250.0, 60.0,
    )
    assert monday.fiber_g == 28.0
    assert monday.expenditure_kcal == 2500.0
    assert monday.trend_weight_kg == 74.2


def test_day_without_food_log_keeps_intake_empty(export_path):
    """Dia só com gasto e tendência: ingestão nula, não zero."""
    tuesday = next(d for d in parse_export(export_path).days if d.date == date(2026, 3, 3))
    assert tuesday.calories_kcal is None and tuesday.protein_g is None
    assert tuesday.expenditure_kcal == 2480.0


def test_targets_come_from_program_in_force(export_path):
    by_date = {d.date: d for d in parse_export(export_path).days}

    # 02/03 é segunda, primeira revisão; 03/03 é terça, ainda a primeira.
    assert by_date[date(2026, 3, 2)].target_calories_kcal == 2200.0
    assert by_date[date(2026, 3, 3)].target_calories_kcal == 2100.0
    # 10/03 é terça e já cai na segunda revisão.
    late = by_date[date(2026, 3, 10)]
    assert (late.target_calories_kcal, late.target_protein_g) == (1850.0, 155.0)


def test_parse_weigh_ins(export_path):
    weigh_ins = parse_export(export_path).weigh_ins
    assert [(w.date, w.weight_kg, w.body_fat_pct) for w in weigh_ins] == [
        (date(2026, 3, 2), 74.5, 15.0),
        (date(2026, 3, 3), 74.0, None),
    ]


def test_parse_rejects_other_spreadsheets(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Planilha1"
    path = tmp_path / "outra.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError):
        parse_export(path)


# ---------------------------------------------------------------------------
# Gravação
# ---------------------------------------------------------------------------


def _import(db, export_path) -> dict:
    return import_export(CatalogStore(db), ATHLETE_ID, parse_export(export_path))


def test_import_is_idempotent(db, export_path):
    _import(db, export_path)
    _import(db, export_path)
    assert db.execute("SELECT count(*) FROM nutrition_daily").fetchone()[0] == 3
    assert db.execute("SELECT count(*) FROM health_metrics").fetchone()[0] == 2


def test_upsert_nutrition_keeps_fields_not_sent(db):
    store = CatalogStore(db)
    day = date(2026, 3, 2)
    store.upsert_nutrition_daily(ATHLETE_ID, day, calories_kcal=2000.0, protein_g=120.0)
    store.upsert_nutrition_daily(ATHLETE_ID, day, expenditure_kcal=2500.0, protein_g=None)

    row = db.execute(
        "SELECT calories_kcal, protein_g, expenditure_kcal FROM nutrition_daily"
    ).fetchone()
    assert row == (2000.0, 120.0, 2500.0)


def test_upsert_nutrition_rejects_unknown_field(db):
    with pytest.raises(ValueError):
        CatalogStore(db).upsert_nutrition_daily(
            ATHLETE_ID, date(2026, 3, 2), sugar_g=10.0
        )


def test_weigh_in_does_not_override_garmin(db, export_path):
    """O Garmin já tinha peso e saúde no dia: nada disso muda."""
    store = CatalogStore(db)
    store.insert_health_daily(
        ATHLETE_ID, date(2026, 3, 2), weight_kg=73.0, hrv_rmssd_ms=40, source="garmin"
    )
    result = _import(db, export_path)

    row = db.execute(
        "SELECT weight_kg, body_fat_pct, hrv_rmssd_ms, source FROM health_metrics "
        "WHERE date = ?",
        [date(2026, 3, 2)],
    ).fetchone()
    assert row == (73.0, 15.0, 40.0, "garmin")
    assert result["weigh_ins"] == 2  # gordura no dia 2, pesagem inteira no dia 3


def test_weigh_in_fills_day_without_health_row(db, export_path):
    _import(db, export_path)
    row = db.execute(
        "SELECT weight_kg, source FROM health_metrics WHERE date = ?",
        [date(2026, 3, 3)],
    ).fetchone()
    assert row == (74.0, "macrofactor")


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------


def test_nutrition_endpoint_window_and_balance(client, db, export_path):
    _import(db, export_path)
    rows = client.get("/nutrition?days=7&end=2026-03-03").json()

    assert [r["date"] for r in rows] == ["2026-03-03", "2026-03-02"]
    assert rows[1]["balance_kcal"] == -500.0
    # Sem comida registrada o balanço é desconhecido, não -2480.
    assert rows[0]["balance_kcal"] is None


def test_nutrition_endpoint_empty(client):
    assert client.get("/nutrition").json() == []


# ---------------------------------------------------------------------------
# Ciclo — o que o Panorama mostra
# ---------------------------------------------------------------------------


def _cycle(db, **kwargs):
    from analytics.nutrition import build_nutrition_cycle

    defaults = dict(first_monday=date(2026, 3, 2), weeks=2, reference=date(2026, 3, 15))
    return build_nutrition_cycle(db, ATHLETE_ID, **{**defaults, **kwargs})


def test_cycle_averages_only_logged_days(db, export_path):
    """Um dia registrado em sete: a média é a dele, não um sétimo."""
    _import(db, export_path)
    week = _cycle(db)["weeks"][0]

    assert week["days_logged"] == 1
    assert week["avg_calories_kcal"] == 2000.0
    assert week["avg_balance_kcal"] == -500.0
    # O gasto existe nos dois dias da semana e entra inteiro na média dele.
    assert week["avg_expenditure_kcal"] == 2490.0
    assert week["reliable"] is False
    assert week["trend_weight_kg"] == 74.1


def test_cycle_week_is_reliable_from_four_days(db):
    store = CatalogStore(db)
    for offset in range(4):
        store.upsert_nutrition_daily(
            ATHLETE_ID, date(2026, 3, 2 + offset), calories_kcal=2000.0,
            expenditure_kcal=2400.0, protein_g=130.0,
        )
    week = _cycle(db)["weeks"][0]
    assert week["days_logged"] == 4 and week["reliable"] is True


def test_cycle_always_returns_every_week(db):
    cycle = _cycle(db)
    assert cycle["has_data"] is False
    assert [w["week_start"] for w in cycle["weeks"]] == ["2026-03-02", "2026-03-09"]
    assert cycle["weeks"][0]["avg_calories_kcal"] is None


def test_zeroed_day_is_not_a_logged_day(export_path):
    """Linha zerada do export não vira dia de 0 kcal."""
    assert date(2026, 3, 11) not in {d.date for d in parse_export(export_path).days}


def test_partial_day_stays_out_of_averages(db):
    """Só o café da manhã anotado: não é um dia de 375 kcal."""
    store = CatalogStore(db)
    store.upsert_nutrition_daily(
        ATHLETE_ID, date(2026, 3, 2), calories_kcal=2000.0, expenditure_kcal=2400.0
    )
    store.upsert_nutrition_daily(
        ATHLETE_ID, date(2026, 3, 3), calories_kcal=375.0, expenditure_kcal=2400.0
    )
    week = _cycle(db)["weeks"][0]

    assert (week["days_logged"], week["days_partial"]) == (1, 1)
    assert week["avg_calories_kcal"] == 2000.0
    assert week["avg_balance_kcal"] == -400.0


def test_nutrition_endpoint_flags_partial_day(client, db):
    CatalogStore(db).upsert_nutrition_daily(
        ATHLETE_ID, date(2026, 3, 3), calories_kcal=375.0, expenditure_kcal=2400.0
    )
    day = client.get("/nutrition?days=1&end=2026-03-03").json()[0]
    assert day["is_partial"] is True and day["balance_kcal"] is None


def test_cycle_summary(db, export_path):
    _import(db, export_path)
    summary = _cycle(db)["summary"]

    assert summary["days_logged"] == 2 and summary["days_total"] == 14
    assert summary["avg_calories_kcal"] == 1900.0
    assert summary["avg_protein_g"] == 130.0
    assert summary["trend_weight_start_kg"] == 74.2
    assert summary["trend_weight_end_kg"] == 74.1
    assert summary["last_logged_on"] == "2026-03-10"


def test_panorama_carries_nutrition(client, db, export_path):
    _import(db, export_path)
    body = client.get("/panorama?weeks=4&zones=false").json()
    assert len(body["nutrition"]["weeks"]) == 4
    assert "summary" in body["nutrition"]
